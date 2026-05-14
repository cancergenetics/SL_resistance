import os
import pandas as pd
import openpyxl

INPUT_FILE = "dev_supp_table_1.xlsx"
HGNC_FILE  = "../../HGNC/hgnc_complete_set.txt"
OUTPUT_PATH = os.path.abspath(
    os.path.join(os.path.dirname(__file__), "../../1_resistance_screens/dev_BRCA1_PARP1_screen.xlsx")
)

SHEET      = "TABLE 1 CRISPR-Cas9 screen"
DRUG_COLS  = {"Olaparib": 2, "Talazoparib": 3, "AZD2461": 4}  # 1-based openpyxl col indices
DATA_START = 2   # first data row (row 1 = header)


def build_hgnc_lookup(hgnc_path):
    hgnc = pd.read_csv(hgnc_path, sep="\t", low_memory=False, dtype=str)
    lookup = {}
    for _, row in hgnc.iterrows():
        hgnc_id = str(row.get("hgnc_id", "")).strip()
        if not hgnc_id or hgnc_id == "nan":
            continue
        tokens = []
        sym = str(row.get("symbol", "")).strip()
        if sym and sym != "nan":
            tokens.append(sym)
        for col in ("prev_symbol", "alias_symbol"):
            raw = str(row.get(col, "")).strip()
            if raw and raw != "nan":
                tokens.extend([t.strip() for t in raw.split("|")])
        for tok in tokens:
            key = tok.upper()
            if key:
                lookup.setdefault(key, hgnc_id)
    print(f"[HGNC] Loaded {len(lookup):,} symbol tokens")
    return lookup


def is_control(gene):
    g = gene.upper()
    if "NONTARGETINGCONTROLGUIDEFORHUMAN" in g.replace("-", "").replace("_", ""):
        return True
    if g.startswith("HSA-MIR") or g.startswith("HSA-LET") or g.startswith("MIR"):
        return True
    return False


def is_highlighted(cell):
    try:
        rgb = cell.fill.fgColor.rgb
        if rgb and rgb not in ("00000000", "FFFFFFFF"):
            return True
    except Exception:
        pass
    try:
        if cell.fill.fill_type and cell.fill.fill_type != "none":
            return True
    except Exception:
        pass
    return False


hgnc_lookup = build_hgnc_lookup(HGNC_FILE)

wb = openpyxl.load_workbook(INPUT_FILE, data_only=True)
ws = wb[SHEET]

# Build pool from Olaparib column (same gene set across all 3 drugs)
all_genes = []
for row in ws.iter_rows(min_row=DATA_START, min_col=2, max_col=2):
    val = row[0].value
    if val is None:
        continue
    gene = str(val).strip()
    if not gene:
        continue
    all_genes.append(gene)

print(f"Total entries in Olaparib column: {len(all_genes)}")

# Filter controls and miRNAs
genes_filtered = [g for g in all_genes if not is_control(g)]
print(f"After control/miRNA filter: {len(genes_filtered)}")

# HGNC filter → pool
pool = [g for g in genes_filtered if g.upper() in hgnc_lookup]
pool_set = set(pool)
print(f"After HGNC filter (pool): {len(pool_set)}")

# Detect highlighted genes per drug
drug_resistance = {}
for drug, col_idx in DRUG_COLS.items():
    highlighted = []
    for row in ws.iter_rows(min_row=DATA_START, min_col=col_idx, max_col=col_idx):
        cell = row[0]
        if cell.value is None:
            continue
        gene = str(cell.value).strip()
        if gene and is_highlighted(cell) and gene in pool_set:
            highlighted.append(gene)
    drug_resistance[drug] = set(highlighted)
    print(f"{drug} highlighted (FDR<0.1): {len(highlighted)} — {sorted(highlighted)}")

# Union resistance
resistance = set().union(*drug_resistance.values())
non_resistance = pool_set - resistance
print(f"\nUnion resistance: {len(resistance)} — {sorted(resistance)}")
print(f"Non-resistance:   {len(non_resistance)}")
print(f"Total:            {len(resistance) + len(non_resistance)}")

out = pd.concat([
    pd.DataFrame({"Class": "Resistance",     "Gene": sorted(resistance)}),
    pd.DataFrame({"Class": "Non-Resistance", "Gene": sorted(non_resistance)}),
], ignore_index=True)

out.to_excel(OUTPUT_PATH, index=False)
print(f"\nSaved: {OUTPUT_PATH}")
